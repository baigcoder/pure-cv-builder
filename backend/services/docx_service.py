"""Editable Word (.docx) export of a CV.

The document mirrors the PDF's content and order in a plain, ATS-friendly
structure: real headings, real bullet lists, no tables or text boxes. The CV's
font and accent colour carry over; Word substitutes the font if it's missing.
"""

import io
import re
from typing import Any

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from services.theme_presets import THEME_PRESETS

DEFAULT_ORDER = ["summary", "experience", "education", "skills", "projects", "publications", "honors", "patents", "talks"]
SECTION_TITLES = {
    "summary": "Summary",
    "experience": "Experience",
    "education": "Education",
    "skills": "Skills",
    "projects": "Projects",
    "publications": "Publications",
    "honors": "Awards",
    "patents": "Patents",
    "talks": "Talks",
}
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def _text(value: Any) -> str:
    if isinstance(value, list):
        return ", ".join(str(item).strip() for item in value if str(item).strip())
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _date(value: str) -> str:
    value = _text(value)
    match = re.fullmatch(r"(\d{4})-(\d{2})(?:-\d{2})?", value)
    if match and 1 <= int(match.group(2)) <= 12:
        return f"{MONTHS[int(match.group(2)) - 1]} {match.group(1)}"
    return "Present" if value.lower() == "present" else value


def _date_range(entry: dict) -> str:
    if _text(entry.get("date")):
        return _date(entry["date"])
    start, end = _date(entry.get("start_date", "")), _date(entry.get("end_date", ""))
    return " – ".join(part for part in (start, end) if part)


def _rgb(hex_color: str | None, fallback: str = "1A1814") -> RGBColor:
    value = (hex_color or "").lstrip("#")
    if not re.fullmatch(r"[0-9a-fA-F]{6}", value):
        value = fallback
    return RGBColor.from_string(value.upper())


def _bottom_border(paragraph, color: RGBColor) -> None:
    """A hairline under section headings, drawn as a paragraph border (not a shape)."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), str(color))
    borders.append(bottom)
    p_pr.append(borders)


class _Writer:
    def __init__(self, font: str, accent: RGBColor, heading_font: str, body_pt: float):
        self.doc = Document()
        self.accent = accent
        self.heading_font = heading_font
        section = self.doc.sections[0]
        section.left_margin = section.right_margin = Cm(1.9)
        section.top_margin = section.bottom_margin = Cm(1.6)
        self.text_width = section.page_width - section.left_margin - section.right_margin

        normal = self.doc.styles["Normal"]
        normal.font.name = font
        normal.font.size = Pt(body_pt)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), font)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.line_spacing = 1.12

    def heading(self, title: str) -> None:
        paragraph = self.doc.add_paragraph(style="Heading 1")
        paragraph.paragraph_format.space_before = Pt(12)
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(title)
        run.font.name = self.heading_font
        run.font.size = Pt(13)
        run.font.bold = False
        run.font.color.rgb = self.accent
        _bottom_border(paragraph, self.accent)

    def entry_line(self, left: str, right: str, *, bold_left: str = "", italic: str = "") -> None:
        """Title on the left, dates on the right via a right-aligned tab stop."""
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(5)
        paragraph.paragraph_format.tab_stops.add_tab_stop(self.text_width, WD_TAB_ALIGNMENT.RIGHT)
        if bold_left:
            paragraph.add_run(bold_left).bold = True
        if left:
            paragraph.add_run(left)
        if italic:
            paragraph.add_run(italic).italic = True
        if right:
            paragraph.add_run("\t" + right)

    def line(self, text: str, *, italic: bool = False, muted: bool = False) -> None:
        if not text:
            return
        run = self.doc.add_paragraph().add_run(text)
        run.italic = italic
        if muted:
            run.font.color.rgb = RGBColor(0x5F, 0x59, 0x4F)

    def bullets(self, items: list[str]) -> None:
        for item in items:
            if _text(item):
                paragraph = self.doc.add_paragraph(_text(item), style="List Bullet")
                paragraph.paragraph_format.space_after = Pt(1)


def build_docx(
    cv_data: dict[str, Any],
    theme: str = "folio",
    design_settings: dict[str, Any] | None = None,
    section_order: list[str] | None = None,
) -> bytes:
    design = design_settings or {}
    preset = THEME_PRESETS.get(theme, {})
    defaults = preset.get("defaults", {})
    font = _text(design.get("fontFamily")) or defaults.get("fontFamily") or "Calibri"
    heading_font = preset.get("design", {}).get("typography", {}).get("font_family", {}).get("name", font)
    if not isinstance(heading_font, str):
        heading_font = font
    accent = _rgb(design.get("primaryColor") or defaults.get("primaryColor"))
    writer = _Writer(font, accent, heading_font, float(design.get("fontSizePt") or 10.5))
    doc = writer.doc

    name = doc.add_paragraph(style="Title")
    name_run = name.add_run(_text(cv_data.get("name")) or "Your Name")
    name_run.font.name = heading_font
    name_run.font.size = Pt(26)
    name_run.font.color.rgb = RGBColor(0x1A, 0x18, 0x14)
    name.paragraph_format.space_after = Pt(2)
    writer.line(_text(cv_data.get("headline")), muted=True)
    contact = [
        _text(cv_data.get(key))
        for key in ("email", "phone", "location", "website")
        if _text(cv_data.get(key))
    ]
    if _text(cv_data.get("linkedin")):
        contact.append(f"linkedin.com/in/{_text(cv_data['linkedin'])}")
    if _text(cv_data.get("github")):
        contact.append(f"github.com/{_text(cv_data['github'])}")
    writer.line("  ·  ".join(contact), muted=True)

    hidden = set(design.get("hiddenSections") or [])
    order = design.get("sectionOrder") or section_order or DEFAULT_ORDER
    order = (["summary"] if "summary" not in order else []) + [s for s in order if s in SECTION_TITLES]
    for section in order:
        if section in hidden:
            continue
        if section == "summary":
            if _text(cv_data.get("summary")):
                writer.heading("Summary")
                writer.line(_text(cv_data["summary"]))
            continue
        entries = [e for e in cv_data.get(section) or [] if isinstance(e, dict) and any(_text(v) for k, v in e.items() if k != "id")]
        if not entries:
            continue
        writer.heading(SECTION_TITLES[section])
        for entry in entries:
            if section == "experience":
                writer.entry_line(
                    f", {_text(entry.get('company'))}" if _text(entry.get("company")) else "",
                    _date_range(entry),
                    bold_left=_text(entry.get("position")) or _text(entry.get("company")),
                )
                writer.line(_text(entry.get("location")), muted=True)
                writer.line(_text(entry.get("summary")))
                writer.bullets(entry.get("highlights") or [])
            elif section == "education":
                degree = " ".join(p for p in (_text(entry.get("degree")), _text(entry.get("area"))) if p)
                writer.entry_line(f", {degree}" if degree else "", _date_range(entry), bold_left=_text(entry.get("institution")))
                writer.line(_text(entry.get("summary")))
                writer.bullets(entry.get("highlights") or [])
            elif section == "skills":
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(2)
                if _text(entry.get("label")):
                    paragraph.add_run(f"{_text(entry['label'])}: ").bold = True
                paragraph.add_run(_text(entry.get("details")))
            elif section == "projects":
                writer.entry_line("", _date_range(entry), bold_left=_text(entry.get("name")))
                writer.line(_text(entry.get("url")), muted=True)
                writer.line(_text(entry.get("summary")))
                writer.bullets(entry.get("highlights") or [])
            elif section == "publications":
                writer.entry_line("", _date(entry.get("date", "")), bold_left=_text(entry.get("title")))
                writer.line(_text(entry.get("authors")), muted=True)
                writer.line(" ".join(p for p in (_text(entry.get("journal")), _text(entry.get("url") or entry.get("doi"))) if p), italic=True)
            else:
                writer.bullets([_text(entry.get("bullet"))])

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
